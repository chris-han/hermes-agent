"""Native scoped evaluation/retention over the incumbent source/definition owners."""

from __future__ import annotations
import csv
import hashlib
import io
import json
from pathlib import Path
import re
from datetime import datetime, timezone

from contracts.etl_pipeline import EtlCommand, EtlReceipt
from contracts.ontology_building import ArtifactPinV1
from eos.artifact_hashing import content_hash
from services.workspace_artifact_service import WorkspaceArtifactService
from storage.workspace_artifact_blob_store import WorkspaceArtifactBlobStore
from services.etl_transformations import evaluate, MAX_ROWS

WORKFLOW_ID = "etl-data-pipeline"
OWNER_REF = "core:native-etl"
MAX_BYTES = 20 * 1024 * 1024


def _digest(data):
    return "sha256:" + hashlib.sha256(data).hexdigest()


def read_saved_recipe(ctx, workflow_hash=None, *, current=True):
    """Verify serialization integrity, consuming TS semantic hashes unchanged."""
    if (
        not ctx.authenticated
        or not ctx.organization_id
        or not ctx.workspace_id
        or not ctx.workspace_root.is_absolute()
    ):
        raise ValueError("AUTHORITY_CONTEXT_REQUIRED")
    scope_bytes = json.dumps(
        [ctx.organization_id, ctx.workspace_id, WORKFLOW_ID, OWNER_REF],
        separators=(",", ":"),
    ).encode()
    directory = (
        ctx.workspace_root
        / ".hermes-workspace/workflow-definition-scopes"
        / hashlib.sha256(scope_bytes).hexdigest()
        / ".hermes-workspace/workflow-definitions"
        / WORKFLOW_ID
    )
    head_path = directory / "head.json"

    def read(path):
        if not path.resolve().is_relative_to(ctx.workspace_root.resolve()):
            raise ValueError("WORKFLOW_SCOPE_MISMATCH")
        if not path.is_file():
            raise ValueError("SAVED_RECIPE_REQUIRED")
        return json.loads(path.read_bytes())

    head = read(head_path)
    if workflow_hash is None:
        workflow_hash = head.get("contentHash")
    if not isinstance(workflow_hash, str) or not re.fullmatch(
        r"sha256:[a-f0-9]{64}", workflow_hash
    ):
        raise ValueError("INVALID_WORKFLOW_HASH")
    if current and head.get("contentHash") != workflow_hash:
        raise ValueError("STALE_WORKFLOW_RECIPE_BASIS")
    state = read(directory / "versions" / (workflow_hash[7:] + ".json"))
    if current and state != head:
        raise ValueError("WORKFLOW_OWNER_VERSION_RECEIPT_MISMATCH")
    if state.get("scope") != {
        "organizationId": ctx.organization_id,
        "workspaceId": ctx.workspace_id,
    }:
        raise ValueError("WORKFLOW_SCOPE_MISMATCH")
    raw = state.get("recipeSnapshot")
    definition = state.get("definition")
    if (
        not isinstance(raw, str)
        or _digest(raw.encode()) != state.get("snapshotBytesHash")
        or json.loads(raw) != definition
    ):
        raise ValueError("WORKFLOW_RECIPE_SNAPSHOT_MISMATCH")
    if (
        state.get("schemaVersion") != "etl_workflow_owner_state.v1"
        or state.get("contentHash") != workflow_hash
        or definition.get("workflowId") != WORKFLOW_ID
        or definition.get("ownerRef") != OWNER_REF
    ):
        raise ValueError("WORKFLOW_OWNER_IDENTITY_MISMATCH")
    return state


def parse_source(data, config, name, *, with_schema=False):
    header = config.get("header_row", 1)
    sheet = config.get("sheet", "Sheet1")
    if (
        type(header) is not int
        or not 1 <= header <= 100
        or not isinstance(sheet, str)
        or not sheet
    ):
        raise ValueError("PARSER_PROFILE_REQUIRED")
    if name.lower().endswith(".csv"):
        encoding = config.get("encoding", "utf-8-sig")
        if encoding not in {"utf-8", "utf-8-sig", "gb18030"}:
            raise ValueError("UNSUPPORTED_ENCODING")
        values = list(csv.reader(io.StringIO(data.decode(encoding))))
    elif name.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            if sheet not in workbook.sheetnames:
                raise ValueError("SHEET_UNAVAILABLE")
            values = []
            for row in workbook[sheet].iter_rows(values_only=True):
                values.append(["" if v is None else str(v) for v in row])
                if len(values) > MAX_ROWS + header:
                    raise ValueError("ETL_ROW_LIMIT_EXCEEDED")
        finally:
            workbook.close()
    else:
        raise ValueError("UNSUPPORTED_SOURCE")
    if len(values) < header:
        raise ValueError("HEADER_UNAVAILABLE")
    columns = values[header - 1]
    if (
        not columns
        or len(columns) != len(set(columns))
        or any(not c or c.startswith("__") for c in columns)
    ):
        raise ValueError("INVALID_ETL_COLUMNS")
    if len(values) - header > MAX_ROWS:
        raise ValueError("ETL_ROW_LIMIT_EXCEEDED")
    rows = []
    for index, row in enumerate(values[header:], header + 1):
        if len(row) != len(columns):
            raise ValueError("SOURCE_ROW_WIDTH_MISMATCH")
        rows.append(
            {
                **dict(zip(columns, row)),
                "__sources": [
                    {"source": config["source"], "sheet": sheet, "row": index}
                ],
            }
        )
    return (rows, columns) if with_schema else rows


def csv_output(rows, columns=None):
    stream = io.StringIO(newline="")
    columns = list(dict.fromkeys(k for row in rows for k in row)) if rows or columns is None else columns
    writer = csv.DictWriter(stream, fieldnames=columns, lineterminator="\n")
    writer.writeheader()
    for row in rows:
        writer.writerow(
            {
                k: json.dumps(v, ensure_ascii=False, separators=(",", ":"))
                if isinstance(v, (list, dict))
                else v
                for k, v in row.items()
            }
        )
    return stream.getvalue()


def standalone_python(
    definition, *, stage_id, workflow_hash=None, source_basis=None, result_columns=None
):
    """Generate only trusted allowlisted evaluator code, never user code."""
    from services import etl_transformations

    clean = json.loads(json.dumps(definition))
    typed = etl_transformations.is_typed_definition(clean)
    if typed:
        if result_columns is None:
            schema = etl_transformations.validate_typed_definition(clean).get(stage_id)
            if schema is None:
                raise ValueError("UNKNOWN_STAGE")
            if not schema["known"]:
                raise ValueError("ETL_STRUCTURAL_SCHEMA_REQUIRED")
            result_columns = [column["name"] for column in schema["columns"]]
        clean = etl_transformations.adapt_etl_definition(clean)
        # This asset executes the preserved v1 evaluator. Binding roles have been
        # compiled into join configs; leaving v2 tags would re-enter typed imports.
        for edge in clean["edges"]:
            edge.pop("config", None)
    for node in clean["nodes"]:
        config = node.get("config", {})
        if config.get("operation") == "source":
            config.pop("artifact_pin", None)
    program = Path(etl_transformations.__file__).read_text()
    # Explicit local file arguments replace governed source refs; no runtime import.
    helpers = (
        "\n"
        + __import__("inspect").getsource(parse_source)
        + "\n"
        + __import__("inspect").getsource(csv_output)
    )
    footer = """
if __name__ == '__main__':
    import argparse
    parser=argparse.ArgumentParser(description='Reproduce the saved native ETL recipe from local input files.')
    parser.add_argument('--input', action='append', required=True, metavar='SOURCE=FILE')
    parser.add_argument('--output', required=True)
    args=parser.parse_args()
    paths=dict(value.split('=',1) for value in args.input)
    source_rows={}; actual_basis=[]
    def load(config,name):
        data=Path(name).read_bytes(); digest='sha256:'+hashlib.sha256(data).hexdigest()
        expected=next((basis for basis in SOURCE_BASIS if basis['source']==config['source']),None)
        if expected and digest!=expected['content_hash']: raise ValueError('ETL_SOURCE_BASIS_MISMATCH')
        rows=parse_source(data,config,name)
        actual_basis.append({'source':config['source'],'content_hash':digest,'total_rows':len(rows)})
        return rows
    for node in RECIPE['nodes']:
        config=node.get('config',{})
        if config.get('operation')=='source':
            name=paths.get(config['source'])
            source_rows[config['source']]=lambda name=name,config=config: load(config,name)
    result=evaluate(RECIPE,source_rows,TARGET,mode='run')
    if result['status']!='completed':
        raise SystemExit('Full checks failed: '+json.dumps(result['issues']))
    output=Path(args.output); output.mkdir(parents=True,exist_ok=True)
    columns=OUTPUT_COLUMNS+['__sources'] if TYPED_RECIPE else None
    (output/'result.csv').write_text(csv_output(result['full_rows'],columns),encoding='utf-8')
    (output/'exceptions.csv').write_text(csv_output(result['exceptions'],columns+['__reasons'] if TYPED_RECIPE else None),encoding='utf-8')
    files={name:'sha256:'+hashlib.sha256((output/name).read_bytes()).hexdigest() for name in ('result.csv','exceptions.csv')}
    (output/'manifest.json').write_text(json.dumps({'workflow_hash':OWNER_HASH,'workflow_version':RECIPE.get('workflowVersion'),'source_basis':actual_basis,'executor_version':result['executor_version'],'stage_id':TARGET,'accepted_rows':result['accepted_rows'],'exceptional_rows':len(result['exceptions']),'files':files},sort_keys=True),encoding='utf-8')
"""
    if stage_id not in {node["id"] for node in clean["nodes"]}:
        raise ValueError("UNKNOWN_STAGE")
    target = stage_id
    return (
        "import csv, io, json, hashlib\nfrom pathlib import Path\n"
        + program
        + helpers
        + "\nRECIPE=json.loads("
        + repr(json.dumps(clean, separators=(",", ":")))
        + ")\nTARGET="
        + repr(target)
        + "\nOWNER_HASH="
        + repr(workflow_hash)
        + "\nSOURCE_BASIS=json.loads("
        + repr(json.dumps(source_basis or [], separators=(",", ":")))
        + ")\nTYPED_RECIPE="
        + repr(typed)
        + "\nOUTPUT_COLUMNS="
        + repr(result_columns if typed else [])
        + "\n"
        + footer
    )


class EtlPipelineService:
    def __init__(self, *, ctx, build_store, authority_resolver):
        self.ctx = ctx
        self.store = build_store
        self.authority = authority_resolver
        self.artifacts = WorkspaceArtifactService(
            build_store=build_store,
            authority_resolver=authority_resolver,
            max_bytes=MAX_BYTES,
        )

    def access(self):
        self.artifacts._access(self.ctx)

    def _source(self, config):
        if not config.get("artifact_pin"):
            raise ValueError("ETL_SOURCE_PIN_REQUIRED")
        pin = ArtifactPinV1.model_validate(config["artifact_pin"])
        artifact = self.artifacts.resolve(ctx=self.ctx, artifact_pin=pin)
        return artifact

    def evaluate(self, command: EtlCommand, *, cancelled=None):
        self.access()
        state = read_saved_recipe(self.ctx, command.workflow_hash)
        definition = state["definition"]
        draft = command.draft_definition
        if draft is not None:
            digest = _digest(json.dumps(draft, ensure_ascii=False, separators=(",", ":")).encode())
            if digest != command.draft_digest:
                raise ValueError("ETL_DRAFT_DIGEST_MISMATCH")
            if any(draft.get(key) != definition.get(key) for key in ("workflowId", "ownerRef", "schemaVersion")):
                raise ValueError("WORKFLOW_OWNER_IDENTITY_MISMATCH")
            if not isinstance(draft.get("nodes"), list) or not isinstance(draft.get("edges"), list) or any(
                not isinstance(node, dict) or not isinstance(node.get("id"), str) or not isinstance(node.get("config"), dict)
                for node in draft["nodes"]
            ) or any(not isinstance(edge, dict) or any(not isinstance(edge.get(key), str) for key in ("source", "target")) for edge in draft["edges"]):
                raise ValueError("INVALID_ETL_DRAFT")
            from services.etl_transformations import adapt_etl_definition, validate_config
            for node in adapt_etl_definition(draft)["nodes"]:
                validate_config(node["config"])
            definition = draft
        basis = []
        source_schemas = {}
        sources = {}
        for node in definition["nodes"]:
            config = node.get("config", {})
            if config.get("operation") != "source":
                continue

            def load(config=config, stage=node["id"]):
                artifact = self._source(config)
                blobs = WorkspaceArtifactBlobStore(ctx=self.ctx, max_bytes=MAX_BYTES)
                with blobs.exact_get(
                    byte_owner_ref=artifact.byte_owner_ref,
                    expected_hash=artifact.content_hash,
                    expected_size=artifact.byte_size,
                ) as stream:
                    rows, columns = parse_source(stream.read(), config, artifact.original_name, with_schema=True)
                source_schemas[stage] = columns
                basis.append(
                    {
                        "source": config["source"],
                        "artifact_pin": config["artifact_pin"],
                        "content_hash": artifact.content_hash,
                        "parser_profile": {
                            k: config.get(k)
                            for k in ("sheet", "header_row", "encoding")
                        },
                        "total_rows": len(rows),
                    }
                )
                return rows

            if config["source"] in sources:
                raise ValueError("DUPLICATE_SOURCE_ALIAS")
            sources[config["source"]] = load
        result = evaluate(
            definition,
            sources,
            command.stage_id,
            mode=command.mode,
            cancelled=cancelled,
            source_schemas=source_schemas,
        )
        if cancelled and cancelled():
            result.update(status="cancelled", full_check_status="not_run")
        # Recheck current authority/revision before returning or atomically publishing.
        self.access()
        read_saved_recipe(self.ctx, command.workflow_hash)
        payload = {
            "workflow_ref": WORKFLOW_ID,
            "workflow_version": definition["workflowVersion"],
            "workflow_hash": state["contentHash"],
            "snapshot_bytes_hash": state["snapshotBytesHash"],
            "source_basis": basis,
            "result": result,
        }
        if draft is not None:
            payload.update(draft=True, draft_digest=command.draft_digest)
        if command.mode != "run":
            payload["result"] = {
                k: v for k, v in result.items() if k not in {"full_rows", "stages"}
            }
            return payload
        exports = {}
        if result["status"] == "completed" and result["full_check_status"] == "passed":
            typed = any(n.get("config", {}).get("operationSchemaVersion") == "etl_operation.v2" for n in definition["nodes"])
            empty_columns = result["columns"] + ["__sources"] if typed else None
            exports = {
                "result.csv": csv_output(result["full_rows"], empty_columns),
                "exceptions.csv": csv_output(result["exceptions"], empty_columns + ["__reasons"] if typed else None),
                "pipeline.py": standalone_python(
                    definition,
                    stage_id=command.stage_id,
                    workflow_hash=state["contentHash"],
                    source_basis=basis,
                    result_columns=result["columns"],
                ),
            }
            if any(n.get("config", {}).get("operationSchemaVersion") == "etl_operation.v2" for n in definition["nodes"]):
                from importlib.util import find_spec
                if find_spec("polars") is None:
                    result["export_capability"] = {"polars": "unavailable", "code": "ETL_POLARS_DEPENDENCY_UNAVAILABLE"}
                else:
                    from services.etl_polars_codegen import generate_polars_export
                    exports.update(generate_polars_export(definition, stage_id=command.stage_id,
                        workflow_hash=state["contentHash"], source_basis=basis))
                    result["export_capability"] = {"polars": "available"}
            manifest = {
                **{k: v for k, v in payload.items() if k != "result"},
                "executor_version": result["executor_version"],
                "stage_id": command.stage_id,
                "accepted_rows": result["accepted_rows"],
                "exceptional_rows": len(result["exceptions"]),
                "files": {
                    name: _digest(data.encode()) for name, data in exports.items()
                },
            }
            exports["manifest.json"] = json.dumps(
                manifest, sort_keys=True, separators=(",", ":")
            )
        if cancelled and cancelled():
            result.update(status="cancelled", full_check_status="not_run")
            exports = {}
        fields = dict(
            schema_version="etl_run_receipt.v1",
            organization_id=self.ctx.organization_id,
            workspace_id=self.ctx.workspace_id,
            actor_id=self.ctx.user_id,
            request_id=command.request_id,
            executor_version="native-etl.v1",
            stage_id=command.stage_id,
            produced_at=datetime.now(timezone.utc).isoformat(),
            exports=exports,
            **payload,
        )
        digest = content_hash(fields)
        receipt = EtlReceipt(
            **fields, receipt_ref="etl-run:" + digest, receipt_hash="sha256:" + digest
        )
        # Compilation can outlive membership or the current saved revision.
        self.access()
        read_saved_recipe(self.ctx, command.workflow_hash)
        self.store.put(receipt)
        return receipt.model_dump(mode="json")

    def read(self, receipt_ref):
        self.access()
        receipt = self.store.get(
            organization_id=self.ctx.organization_id,
            workspace_id=self.ctx.workspace_id,
            artifact_ref=receipt_ref,
            contract_type=EtlReceipt,
        )
        if receipt is None:
            raise ValueError("ETL_RECEIPT_UNAVAILABLE")
        state = read_saved_recipe(self.ctx, receipt.workflow_hash, current=False)
        if state["snapshotBytesHash"] != receipt.snapshot_bytes_hash:
            raise ValueError("WORKFLOW_RECIPE_SNAPSHOT_MISMATCH")
        used = {item["source"]: item for item in receipt.source_basis}
        for node in state["definition"]["nodes"]:
            config = node.get("config", {})
            if config.get("operation") == "source" and config.get("source") in used:
                artifact = self._source(config)
                basis = used.pop(config["source"])
                if (
                    basis["artifact_pin"] != config["artifact_pin"]
                    or basis["content_hash"] != artifact.content_hash
                ):
                    raise ValueError("ETL_SOURCE_BASIS_MISMATCH")
        if used:
            raise ValueError("ETL_SOURCE_BASIS_MISMATCH")
        current = read_saved_recipe(self.ctx)["contentHash"] == receipt.workflow_hash
        return receipt, current

    def source_profile(self, command):
        self.access()
        artifact = self.artifacts.resolve(
            ctx=self.ctx,
            artifact_pin=ArtifactPinV1.model_validate(command.artifact_pin),
        )
        blobs = WorkspaceArtifactBlobStore(ctx=self.ctx, max_bytes=MAX_BYTES)
        with blobs.exact_get(
            byte_owner_ref=artifact.byte_owner_ref,
            expected_hash=artifact.content_hash,
            expected_size=artifact.byte_size,
        ) as stream:
            data = stream.read()
        sheets = ["Sheet1"]
        if artifact.original_name.lower().endswith(".xlsx"):
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                sheets = workbook.sheetnames
            finally:
                workbook.close()
        sheet = command.sheet or (sheets[0] if len(sheets) == 1 else None)
        if sheet is None:
            self.access()
            return {
                "sheets": sheets,
                "requires_sheet": True,
                "artifact_pin": command.artifact_pin,
            }
        config = {
            "source": "input",
            "sheet": sheet,
            "header_row": command.header_row,
            "encoding": command.encoding,
        }
        rows, columns = parse_source(data, config, artifact.original_name, with_schema=True)
        self.access()
        return {
            "artifact_pin": command.artifact_pin,
            "original_name": artifact.original_name,
            "sheets": sheets,
            "sheet": sheet,
            "header_row": command.header_row,
            "encoding": command.encoding,
            "total_rows": len(rows),
            "columns": columns,
            "unique_columns": [
                c for c in columns if len({str(r[c]) for r in rows}) == len(rows)
            ],
            "rows": rows[:20],
        }

    def source_row(self, command):
        self.access()
        state = read_saved_recipe(self.ctx, command.workflow_hash, current=False)
        matches = [
            n["config"]
            for n in state["definition"]["nodes"]
            if n.get("config", {}).get("operation") == "source"
            and n["config"].get("source") == command.source
        ]
        if len(matches) != 1:
            raise ValueError("SOURCE_UNAVAILABLE")
        config = matches[0]
        if config.get("sheet", "Sheet1") != command.sheet:
            raise ValueError("SOURCE_LOCATION_MISMATCH")
        artifact = self._source(config)
        blobs = WorkspaceArtifactBlobStore(ctx=self.ctx, max_bytes=MAX_BYTES)
        with blobs.exact_get(
            byte_owner_ref=artifact.byte_owner_ref,
            expected_hash=artifact.content_hash,
            expected_size=artifact.byte_size,
        ) as stream:
            rows = parse_source(stream.read(), config, artifact.original_name)
        found = [r for r in rows if r["__sources"][0]["row"] == command.row]
        if len(found) != 1:
            raise ValueError("SOURCE_ROW_UNAVAILABLE")
        self.access()
        return {
            "artifact_pin": config["artifact_pin"],
            "content_hash": artifact.content_hash,
            "original_name": artifact.original_name,
            "location": found[0]["__sources"][0],
            "row": found[0],
        }

    def list_runs(self):
        self.access()
        items = []
        for pin in self.store.list_artifact_pins(
            organization_id=self.ctx.organization_id,
            workspace_id=self.ctx.workspace_id,
            schema_version="etl_run_receipt.v1",
        ):
            receipt, current = self.read(pin.artifact_ref)
            items.append(
                {
                    "receipt_ref": receipt.receipt_ref,
                    "receipt_hash": receipt.receipt_hash,
                    "workflow_hash": receipt.workflow_hash,
                    "workflow_version": receipt.workflow_version,
                    "produced_at": receipt.produced_at,
                    "status": receipt.result["status"],
                    "is_current": current,
                }
            )
        return sorted(items, key=lambda item: item["produced_at"], reverse=True)
